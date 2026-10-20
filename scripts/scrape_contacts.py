#!/usr/bin/env python3
"""批量抓取网站公开联系方式。无需启动原项目的前后端服务。"""
import argparse
import csv
import heapq
import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urljoin, urlsplit, urlunsplit

try:
    import requests
    import phonenumbers
    from contact_core.html_text import build_page, page_meta
    from contact_core.extract_emails import extract_emails
    from contact_core.extract_phones import extract_phones, region_from_domain
    from contact_core.extract_socials import extract_socials
except ImportError as exc:
    raise SystemExit(f"缺少依赖：{exc}\n请先运行：python -m pip install -r requirements.txt")

SKIP = re.compile(r'\.(?:pdf|png|jpe?g|gif|svg|webp|ico|css|js|json|xml|zip|gz|mp[34]|woff2?)(?:$|\?)', re.I)
NOISE = re.compile(r'/(?:login|logout|cart|checkout|wp-admin|search|register)(?:/|$)', re.I)
TRACKING = re.compile(r'^(?:utm_.*|fbclid|gclid|msclkid)$', re.I)
SUMMARY = ['输入网址', '网站', '公司名称或网站标题', '邮箱', '电话', '社交账号', '地址', '成功页数', '状态', '备注', '抓取时间']
DETAIL = ['输入网址', '公司名称或网站标题', '类型', '联系方式', '来源页面']
LOG = ['输入网址', '页面', '结果', '说明']
from directory_scraper import EXTRA_HEADERS, directory_kind, scrape_directory
SUMMARY += EXTRA_HEADERS


def normalize_url(value):
    value = value.strip()
    if not value:
        raise ValueError('网址为空')
    if '://' not in value:
        value = 'https://' + value
    p = urlsplit(value)
    if p.scheme not in ('http', 'https') or not p.hostname or p.username or p.password:
        raise ValueError('请输入有效的 HTTP/HTTPS 网址')
    if re.search(r'\s', p.netloc):
        raise ValueError('网址域名不能包含空格')
    _ = p.port  # Validate malformed port numbers.
    return urlunsplit((p.scheme.lower(), p.netloc.lower(), p.path or '/',
                       urlencode([(k, v) for k, v in parse_qsl(p.query, keep_blank_values=True)
                                  if not TRACKING.match(k)]), ''))


def site_host(url):
    host = urlsplit(url).hostname or ''
    return host[4:] if host.startswith('www.') else host


def priority(url, label):
    value = urlsplit(url).path.lower() + ' ' + label.lower()
    if re.search(r'contact|kontakt|impressum|联系|联络', value):
        return 100
    if re.search(r'about|company|team|关于|公司', value):
        return 80
    if re.search(r'support|location|office|支持|地址', value):
        return 60
    return 10


def fetch(session, url, timeout, allowed_host=None):
    """Limit body size and follow redirects explicitly to keep inner pages on site."""
    for _ in range(6):
        with session.get(url, timeout=timeout, stream=True, allow_redirects=False) as response:
            if response.is_redirect or response.is_permanent_redirect:
                target = normalize_url(urljoin(url, response.headers['Location']))
                if allowed_host and site_host(target) != allowed_host:
                    raise ValueError('跳转到站外页面，已跳过')
                url = target
                continue
            response.raise_for_status()
            content_type = response.headers.get('Content-Type', '').lower()
            if content_type and 'html' not in content_type:
                raise ValueError('非 HTML 页面，已跳过')
            data = bytearray()
            for chunk in response.iter_content(65536):
                data.extend(chunk)
                if len(data) > 3_000_000:
                    raise ValueError('页面超过 3 MB，已跳过')
            # Let the HTML parser detect meta charset for sites without an HTTP charset.
            payload = bytes(data)
            if 'charset=' in content_type:
                payload = payload.decode(response.encoding or 'utf-8', errors='replace')
            page = build_page(url, payload)
            # Social extractors expect a decoded string.
            page.html = str(page.soup)
            return page
    raise ValueError('跳转次数过多')


def structured_info(page):
    names, addresses = [], []

    def visit(item):
        if isinstance(item, list):
            for child in item:
                visit(child)
        elif isinstance(item, dict):
            kind = item.get('@type', [])
            kinds = kind if isinstance(kind, list) else [kind]
            if any(str(k) in ('Organization', 'Corporation', 'LocalBusiness', 'ProfessionalService', 'Store') for k in kinds):
                if isinstance(item.get('name'), str):
                    names.append(item['name'])
                address = item.get('address')
                if isinstance(address, str):
                    addresses.append(address)
                elif isinstance(address, dict):
                    addresses.append(' '.join(str(address[k]) for k in
                        ('streetAddress', 'addressLocality', 'addressRegion', 'postalCode', 'addressCountry')
                        if isinstance(address.get(k), (str, int))))
            for child in item.values():
                if isinstance(child, (dict, list)):
                    visit(child)

    for tag in page.soup.find_all('script', type='application/ld+json'):
        try:
            visit(json.loads(tag.string or tag.get_text()))
        except (ValueError, RecursionError):
            pass
    addresses.extend(tag.get_text(' ', strip=True) for tag in page.soup.find_all('address'))
    return names, [a for a in addresses if a]


def crawl(raw, args, session):
    record = dict.fromkeys(SUMMARY, '')
    record.update({'输入网址': raw, '成功页数': 0, '抓取时间': datetime.now().isoformat(timespec='seconds')})
    hits, logs = {}, []
    try:
        seed = normalize_url(raw)
    except ValueError as exc:
        record.update({'状态': '网址无效', '备注': str(exc)})
        return record, [], [[raw, raw, '失败', str(exc)]]
    record['网站'] = seed
    queue = [(0, 0, seed, 0)]
    seen, visited = {seed}, set()
    host, sequence, attempts = site_host(seed), 0, 0
    started = time.monotonic()

    def add(kind, value, source):
        key = (kind, value)
        if key not in hits:
            hits[key] = []
        if source not in hits[key]:
            hits[key].append(source)

    while queue and attempts < args.max_pages:
        if time.monotonic() - started >= args.site_timeout:
            logs.append([raw, seed, '停止', '达到单站时间预算'])
            break
        _, _, url, depth = heapq.heappop(queue)
        if url in visited:
            continue
        visited.add(url)
        attempts += 1
        if attempts > 1:
            time.sleep(args.delay)
        try:
            page = fetch(session, url, args.timeout, host if depth else None)
            final = normalize_url(page.url)
            if depth:
                if final != url and final in visited:
                    continue
            else:
                host = site_host(final)
                record['网站'] = final
            visited.add(final)
            names, addresses = structured_info(page)
            if not record['公司名称或网站标题']:
                record['公司名称或网站标题'] = (names or [page_meta(page.soup)[0] or host])[0]
            social, wa = extract_socials(page)
            region = args.region or region_from_domain(host)
            for item in extract_emails(page):
                add('邮箱', item['value'], page.url)
            for item in extract_phones(page, region=region, extra_candidates=wa):
                add('电话', item['value'], page.url)
            for item in social:
                add(item['platform'], item['value'], page.url)
            for address in addresses:
                add('地址', address, page.url)
            record['成功页数'] += 1
            logs.append([raw, page.url, '成功', ''])
            print(f"  [{attempts}/{args.max_pages}] {page.url}", flush=True)
            if depth >= args.depth:
                continue
            for anchor in page.anchors:
                if not anchor['resolved']:
                    continue
                try:
                    target = normalize_url(anchor['resolved'])
                except ValueError:
                    continue
                if (target in seen or site_host(target) != host or SKIP.search(target)
                        or NOISE.search(urlsplit(target).path) or len(seen) >= 2000):
                    continue
                seen.add(target)
                sequence += 1
                heapq.heappush(queue, (-priority(target, anchor['text']), sequence, target, depth + 1))
        except (requests.RequestException, ValueError, UnicodeError) as exc:
            logs.append([raw, url, '失败', str(exc)[:500]])
            print(f"  跳过：{url}：{exc}", flush=True)
    for kind in ('邮箱', '电话', '地址'):
        record[kind] = '\n'.join(value for (k, value) in hits if k == kind)
    record['社交账号'] = '\n'.join(f'{k}: {v}' for k, v in hits if k not in ('邮箱', '电话', '地址'))
    failures = [row for row in logs if row[2] != '成功']
    record['状态'] = ('部分成功' if failures else '成功') if record['成功页数'] else '失败'
    notes = []
    if not hits and record['成功页数']:
        notes.append('已访问页面未检出联系方式；不代表网站没有联系方式')
    if queue and attempts >= args.max_pages:
        notes.append('已达到页面上限')
    notes.extend(row[3] for row in failures[:3])
    record['备注'] = '；'.join(notes)
    details = [[raw, record['公司名称或网站标题'], k, v, source]
               for (k, v), sources in hits.items() for source in sources]
    return record, details, logs


def safe_csv(value):
    text = str(value)
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) or text.startswith(('\t', '\r')) else value


def export_tables(folder, records, details, logs, excel=True):
    tables = [('公司汇总', SUMMARY, [[r[h] for h in SUMMARY] for r in records]),
              ('联系方式明细', DETAIL, details), ('抓取日志', LOG, logs)]
    folder.mkdir(parents=True, exist_ok=True)
    for name, headers, rows in tables:
        target = folder / f'{name}.csv'
        temp = target.with_suffix('.csv.tmp')
        with temp.open('w', encoding='utf-8-sig', newline='') as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows([[safe_csv(v) for v in row] for row in rows])
        temp.replace(target)
    if excel:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        wb = Workbook()
        wb.remove(wb.active)
        for name, headers, rows in tables:
            ws = wb.create_sheet(name)
            for row in [headers] + rows:
                ws.append([ILLEGAL_CHARACTERS_RE.sub('', v) if isinstance(v, str) else v for v in row])
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'  # URLs, +phones and scraped strings are never formulas.
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='245A81')
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = 32
            ws.row_dimensions[1].height = 28
        temp = folder / '联系方式.tmp.xlsx'
        wb.save(temp)
        temp.replace(folder / '联系方式.xlsx')


def positive_int(value):
    number = int(value)
    if number <= 0:
        raise argparse.ArgumentTypeError('必须大于 0')
    return number


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('urls', nargs='*', help='一个或多个网址')
    parser.add_argument('-i', '--input', type=Path, help='UTF-8 文本文件，每行一个网址')
    parser.add_argument('-o', '--output', type=Path, default=Path.cwd() / 'output' / 'website-contacts',
                        help='输出根目录，默认当前工作目录/output/website-contacts')
    parser.add_argument('--max-pages', type=positive_int, default=20, help='每站最多尝试页面数，默认20')
    parser.add_argument('--depth', type=int, default=2, help='链接深度，默认2')
    parser.add_argument('--timeout', type=positive_int, default=15, help='单次网络连接/读取超时秒数')
    parser.add_argument('--site-timeout', type=positive_int, default=120, help='单站调度时间预算秒数')
    parser.add_argument('--delay', type=float, default=0.5, help='页面间隔秒数')
    parser.add_argument('--region', type=str.upper, help='无国际区号电话的国家，例如 CN、US、GB')
    parser.add_argument('--csv-only', action='store_true', help='仅输出CSV，不需要openpyxl')
    parser.add_argument('--generic', action='store_true', help='禁用展会/会员目录自动识别，按普通网站抓取')
    parser.add_argument('--directory-pages', type=positive_int, default=200, help='目录分页上限，默认200；与普通网站max-pages独立')
    parser.add_argument('--directory-workers', type=positive_int, default=3, help='目录读取并发数，默认3，最大8')
    parser.add_argument('--refresh-directories', action='store_true', help='重新读取目录，忽略24小时缓存')
    args = parser.parse_args()
    if args.depth < 0 or not 0 <= args.delay <= 60:
        parser.error('depth 必须非负；delay 必须介于 0 和 60 秒')
    if args.region and args.region not in phonenumbers.SUPPORTED_REGIONS:
        parser.error('未知的电话国家代码')
    if args.directory_workers > 8:
        parser.error('目录并发数不能大于8')
    urls = list(args.urls)
    if args.input:
        try:
            urls += [line.strip() for line in args.input.read_text(encoding='utf-8-sig').splitlines()
                     if line.strip() and not line.lstrip().startswith('#')]
        except (OSError, UnicodeError) as exc:
            parser.error(f'无法读取网址文件：{exc}')
    if not urls:
        urls = input('请输入网址（多个网址用空格隔开）：').split()
    urls = list(dict.fromkeys(urls))
    if not urls:
        parser.error('没有提供网址')
    if not args.csv_only:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            parser.error('Excel导出需要openpyxl，请安装依赖或添加 --csv-only')
    folder = args.output.resolve() / datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    records, details, logs = [], [], []
    active_directory = None
    session = requests.Session()
    session.headers['User-Agent'] = 'CompanyContactScraper/1.0'
    try:
        for index, url in enumerate(urls, 1):
            print(f'\n网站 {index}/{len(urls)}：{url}', flush=True)
            if not args.generic and directory_kind(url):
                def checkpoint(batch, items, events):
                    nonlocal active_directory
                    active_directory = (list(batch), list(items), list(events))
                    export_tables(folder, records + batch, details + items, logs + events, excel=False)
                batch, items, events = scrape_directory(url, args, checkpoint)
                records.extend(batch)
                details.extend(items)
                logs.extend(events)
                active_directory = None
                export_tables(folder, records, details, logs, excel=False)
                continue
            row, items, events = crawl(url, args, session)
            records.append(row)
            details.extend(items)
            logs.extend(events)
            # Persist after each completed site, including failed sites.
            export_tables(folder, records, details, logs, excel=False)
    except KeyboardInterrupt:
        if active_directory:
            batch, items, events = active_directory
            for row in batch:
                row['备注'] = '；'.join(filter(None, [row['备注'], '用户中断，目录不完整']))
            records.extend(batch)
            details.extend(items)
            logs.extend(events)
            logs.append([url, url, '目录不完整', '用户中断，已保存完成的目录页'])
        print('\n已中断，正在保存已完成的网站或目录页。')
    finally:
        session.close()
        if records or logs:
            export_tables(folder, records, details, logs, excel=not args.csv_only)
            print(f'\n已保存 {len(records)} 条公司/网站记录：{folder}')
    directory_failed = any(row[2] in ('目录失败', '目录不完整') for row in logs)
    return 0 if records and any(r['成功页数'] for r in records) and not directory_failed else 1


if __name__ == '__main__':
    sys.exit(main())
