import argparse
import csv
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import requests
from openpyxl import load_workbook
from scrape_contacts import crawl, export_tables, normalize_url


class Handler(BaseHTTPRequestHandler):
    paths = []

    def log_message(self, *args):
        pass

    def do_GET(self):
        self.paths.append(self.path)
        if self.path == '/bad':
            self.send_error(500)
            return
        pages = {
            '/': '<title>测试公司</title><a href="/product">产品</a><a href="/contact">联系我们</a><a href="mailto:sales@realbusiness.cn">邮件</a><a href="https://other.invalid/contact">外站</a><a href="/bad">失败页</a>',
            '/contact': '<p>sales@realbusiness.cn</p><p>service [at] realbusiness [dot] cn</p><a href="tel:+442079460958">电话</a><a href="https://www.linkedin.com/company/realbusiness">LinkedIn</a><address>上海市测试路100号</address><a href="/contact#top">自身</a><a href="/?utm_source=test">首页</a>',
            '/product': '<p>产品页</p>',
        }
        self.send_response(200)
        self.send_header('Content-Type', 'text/html; charset=utf-8')
        self.end_headers()
        self.wfile.write(pages.get(self.path, '').encode('utf-8'))


class ScraperTest(unittest.TestCase):
    def test_crawl_and_export(self):
        Handler.paths = []
        server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        try:
            args = argparse.Namespace(max_pages=8, depth=2, timeout=2, site_timeout=20, delay=0, region='CN')
            with requests.Session() as session:
                session.trust_env = False
                record, details, logs = crawl(f'http://127.0.0.1:{server.server_port}', args, session)
            self.assertEqual(record['公司名称或网站标题'], '测试公司')
            self.assertEqual(record['成功页数'], 3)
            self.assertEqual(record['状态'], '部分成功')
            self.assertEqual(Handler.paths[1], '/contact')
            self.assertEqual(len(Handler.paths), 4)
            self.assertEqual(record['邮箱'].count('sales@realbusiness.cn'), 1)
            self.assertIn('service@realbusiness.cn', record['邮箱'])
            self.assertIn('+442079460958', record['电话'])
            self.assertIn('上海市', record['地址'])
            self.assertIn('linkedin.com/company/realbusiness', record['社交账号'])
            self.assertEqual(sum(r[3] == 'sales@realbusiness.cn' for r in details), 2)
            # Verify formula-like source text and international phone numbers remain text.
            record['公司名称或网站标题'] = '=1+1'
            with tempfile.TemporaryDirectory() as temp:
                folder = Path(temp)
                export_tables(folder, [record], details, logs)
                wb = load_workbook(folder / '联系方式.xlsx')
                self.assertEqual(wb.sheetnames, ['公司汇总', '联系方式明细', '抓取日志'])
                self.assertEqual(wb['公司汇总']['C2'].data_type, 's')
                self.assertEqual(wb['公司汇总']['E2'].value, '+442079460958')
                self.assertEqual(wb['公司汇总'].freeze_panes, 'A2')
                wb.close()
                with (folder / '公司汇总.csv').open(encoding='utf-8-sig', newline='') as f:
                    rows = list(csv.reader(f))
                self.assertEqual(rows[1][2], "'=1+1")
                self.assertIn('sales@realbusiness.cn', rows[1][3])
        finally:
            server.shutdown()
            server.server_close()

    def test_urls(self):
        self.assertEqual(normalize_url('www.realbusiness.cn?a=1&utm_source=x#foo'), 'https://www.realbusiness.cn/?a=1')
        for value in ['ftp://a.cn', 'https://', 'not a url', 'http://a.cn:bad']:
            with self.assertRaises(ValueError):
                normalize_url(value)

    def test_invalid_site_kept(self):
        record, details, logs = crawl('https://', None, None)
        self.assertEqual(record['状态'], '网址无效')
        self.assertEqual(details, [])
        self.assertEqual(len(logs), 1)


if __name__ == '__main__':
    unittest.main()
